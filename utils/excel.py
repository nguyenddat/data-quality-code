from pathlib import Path
from typing import Any

from openpyxl.styles import Font
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from core.settings import config

SCHEMA_HEADERS = ["Schema", "Table Count", "Total Size", "DB Comment"]


def load_or_create_workbook() -> Any:
    """Open the database workbook or create a new workbook."""
    result_file = Path(config.db_result_file)
    result_file.parent.mkdir(parents=True, exist_ok=True)
    return load_workbook(result_file) if result_file.exists() else Workbook()


def get_or_create_sheet(workbook: Any, name: str) -> Any:
    """Return a named sheet, reusing the default empty sheet when possible."""
    if name in workbook.sheetnames:
        return workbook[name]

    active = workbook.active
    if (
        len(workbook.sheetnames) == 1
        and active.title == "Sheet"
        and active.max_row == 1
        and active.max_column == 1
        and active.cell(1, 1).value is None
    ):
        active.title = name
        return active

    return workbook.create_sheet(name)


def setup_schema_sheet(workbook: Any) -> tuple[Any, list[dict[str, Any]]]:
    """Load or create the Schemas sheet and return its existing data."""
    sheet = get_or_create_sheet(workbook, "Schemas")

    for column, header in enumerate(SCHEMA_HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)

    data = [
        {
            "schema": sheet.cell(row=row, column=1).value,
            "table_count": sheet.cell(row=row, column=2).value,
            "total_size": sheet.cell(row=row, column=3).value,
            "comment": sheet.cell(row=row, column=4).value,
        }
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row=row, column=1).value
    ]

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{max(sheet.max_row, 1)}"
    return sheet, data


def setup_table_sheet() -> tuple[Any, Any]:
    headers = [
        "Schema",
        "Table Name",
        "Phân loại",
        "Comment",
        "Row Count",
        "Columns",
        "Total Size",
        "Data Size",
        "Index Size",
        "Has PK",
        "FK Count",
    ]

    result_file = Path(config.db_result_file)
    workbook = load_or_create_workbook()
    sheet = get_or_create_sheet(workbook, "Tables")

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)

    catalog = [
        ("Tier", "Ý nghĩa"),
        (1, "Bảng nghiệp vụ"),
        (2, "Bảng hỗ trợ"),
        (3, "Bảng kỹ thuật"),
    ]
    for row, values in enumerate(catalog, start=1):
        for column, value in enumerate(values, start=13):
            cell = sheet.cell(row=row, column=column, value=value)
            if row == 1:
                cell.font = Font(bold=True)

    if "TierCatalog" in sheet.tables:
        del sheet.tables["TierCatalog"]
    tier_table = Table(displayName="TierCatalog", ref="M1:N4")
    tier_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(tier_table)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:K1"
    workbook.save(result_file)
    return workbook, sheet


def setup_column_sheet(workbook: Any) -> Any:
    headers = [
        "Schema",
        "Table",
        "Column",
        "Data Type",
        "Full Type",
        "Nullable",
        "Default",
        "PK",
        "FK",
        "UQ",
        "PII/PHI",
        "Null pct",
        "Distinct Count",
        "Comment",
        "Null Warning",
        "Dirty Check",
        "NotNullable",
        "Null Classify",
    ]
    sheet = get_or_create_sheet(workbook, "Columns")
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:Q{max(sheet.max_row, 1)}"
    return sheet


def save_excel(workbook: Any) -> Path:
    result_file = Path(config.db_result_file)
    result_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(result_file)
    return result_file


def load_or_create_consistency_workbook() -> Any:
    """Open the separate workbook used for Consistency results."""
    result_file = Path(config.db_consistency_result_file)
    result_file.parent.mkdir(parents=True, exist_ok=True)
    return load_workbook(result_file) if result_file.exists() else Workbook()


def setup_consistency_sheet(
    workbook: Any,
    name: str,
    headers: list[str],
    rows: list[list[Any]],
) -> Any:
    """Create/reset one Consistency step sheet and write its tabular results."""
    sheet = get_or_create_sheet(workbook, name)
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)

    for row_number, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_number, column=column, value=value)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(sheet.max_row, 1)}"
    return sheet


def save_consistency_excel(workbook: Any) -> Path:
    """Save the separate Consistency workbook."""
    result_file = Path(config.db_consistency_result_file)
    result_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(result_file)
    return result_file


def load_or_create_uniqueness_workbook() -> Any:
    """Open the separate workbook used for Uniqueness results."""
    result_file = Path(config.db_uniqueness_result_file)
    result_file.parent.mkdir(parents=True, exist_ok=True)
    return load_workbook(result_file) if result_file.exists() else Workbook()


def setup_uniqueness_sheet(
    workbook: Any,
    name: str,
    headers: list[str],
    rows: list[list[Any]],
) -> Any:
    """Create/reset one Uniqueness result sheet."""
    sheet = get_or_create_sheet(workbook, name)
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)

    for row_number, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_number, column=column, value=value)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(sheet.max_row, 1)}"
    return sheet


def save_uniqueness_excel(workbook: Any) -> Path:
    """Save the separate Uniqueness workbook."""
    result_file = Path(config.db_uniqueness_result_file)
    result_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(result_file)
    return result_file
